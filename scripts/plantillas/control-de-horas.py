"""
Genera public/plantillas/control-de-horas-trabajadas.xlsx, la planilla gratis
que se descarga desde /blog/control-de-horas-trabajadas.

Correr con:  python3 scripts/plantillas/control-de-horas.py
Necesita openpyxl (pip install openpyxl). Se regenera entera: no editar el
.xlsx a mano, cambiar acá y volver a correr.

Qué trae:
- Hoja "Horas": una fila por bloque de trabajo (fecha, cliente, tarea,
  entrada, salida, descanso, facturable) y las fórmulas que calculan las
  horas en hh:mm y en decimal, la tarifa del cliente y el importe.
- Hoja "Clientes": cliente, tarifa por hora y moneda. La hoja Horas la lee
  con VLOOKUP.
- Hoja "Resumen": horas, horas facturables e importe por cliente, totales y
  un bloque por rango de fechas.
- Hoja "Instrucciones".

Las fórmulas usan solo funciones que existen igual en Excel, Google Sheets,
LibreOffice y Numbers (IF, OR, MOD, ROUND, IFERROR, VLOOKUP, SUMIF, SUMIFS).
El ejemplo precargado es la misma semana que cuenta el artículo.
"""

from datetime import date, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "public" / "plantillas" / "control-de-horas-trabajadas.xlsx"

# Filas con fórmulas listas para completar (además del ejemplo).
ROWS = 400
# Filas de la hoja Clientes que lee el resumen.
CLIENT_ROWS = 30

INDIGO = "4F46E5"
SLATE_50 = "F8FAFC"
SLATE_200 = "E2E8F0"

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor=INDIGO)
formula_fill = PatternFill("solid", fgColor=SLATE_50)
thin = Side(style="thin", color=SLATE_200)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
title_font = Font(bold=True, size=14)
muted = Font(color="64748B")


def style_header(ws, row, count):
    for col in range(1, count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ----------------------------------------------------------------------------
# Clientes
# ----------------------------------------------------------------------------
wb = Workbook()
clients = wb.active
clients.title = "Clientes"
clients.append(["Cliente", "Tarifa por hora", "Moneda"])
style_header(clients, 1, 3)
EXAMPLE_CLIENTS = [
    ("Estudio Galley", 40, "USD"),
    ("Baratie Labs", 45, "USD"),
    ("All Blue Ventures", 60, "USD"),
    ("Mi negocio", 0, "USD"),
]
for row in EXAMPLE_CLIENTS:
    clients.append(list(row))
for r in range(2, CLIENT_ROWS + 2):
    clients.cell(row=r, column=2).number_format = "#,##0.00"
    for c in range(1, 4):
        clients.cell(row=r, column=c).border = border
clients.freeze_panes = "A2"
widths(clients, [28, 16, 10])
clients["E2"] = "Cada cliente con su tarifa. La hoja Horas la busca sola por el nombre."
clients["E3"] = "“Mi negocio” es para las horas propias (propuestas, administración): tarifa 0 y no facturable."
clients["E2"].font = muted
clients["E3"].font = muted

# ----------------------------------------------------------------------------
# Horas
# ----------------------------------------------------------------------------
ws = wb.create_sheet("Horas", 0)
HEADERS = [
    "Fecha",
    "Cliente",
    "Tarea",
    "Entrada",
    "Salida",
    "Descanso",
    "Horas",
    "Horas (decimal)",
    "Facturable",
    "Tarifa por hora",
    "Importe",
]
ws.append(HEADERS)
style_header(ws, 1, len(HEADERS))
ws.row_dimensions[1].height = 30

# La misma semana del artículo: 23:45 en total, 22:45 facturables.
EXAMPLE = [
    (date(2026, 10, 19), "Estudio Galley", "Auditoría de producto: revisión del onboarding", time(9, 0), time(12, 30), 0, "Sí"),
    (date(2026, 10, 19), "Baratie Labs", "Discovery: entrevistas a usuarios (2)", time(14, 0), time(16, 0), 0, "Sí"),
    (date(2026, 10, 19), "All Blue Ventures", "Sesión de mentoría", time(17, 0), time(17, 45), 0, "Sí"),
    (date(2026, 10, 20), "Estudio Galley", "Auditoría: informe de hallazgos", time(9, 0), time(12, 0), 0, "Sí"),
    (date(2026, 10, 20), "Baratie Labs", "Roadmap Q3: priorización", time(15, 0), time(16, 30), 0, "Sí"),
    (date(2026, 10, 21), "Estudio Galley", "Presentación de la auditoría al equipo", time(10, 0), time(11, 30), 0, "Sí"),
    (date(2026, 10, 21), "Baratie Labs", "Discovery: síntesis y oportunidades", time(14, 0), time(18, 0), 30, "Sí"),
    (date(2026, 10, 22), "Estudio Galley", "Propuesta de rediseño del checkout", time(9, 0), time(11, 45), 0, "Sí"),
    (date(2026, 10, 22), "Mi negocio", "Propuesta comercial para un cliente nuevo", time(16, 0), time(17, 0), 0, "No"),
    (date(2026, 10, 23), "Baratie Labs", "Roadmap Q3: documento final", time(9, 0), time(12, 15), 0, "Sí"),
    (date(2026, 10, 23), "All Blue Ventures", "Sesión de mentoría", time(15, 0), time(16, 0), 0, "Sí"),
]
for i, (d, client, task, start, end, brk, billable) in enumerate(EXAMPLE, start=2):
    ws.cell(row=i, column=1, value=d)
    ws.cell(row=i, column=2, value=client)
    ws.cell(row=i, column=3, value=task)
    ws.cell(row=i, column=4, value=start)
    ws.cell(row=i, column=5, value=end)
    # Descanso en minutos convertido a fracción de día, que es como Excel guarda las horas.
    ws.cell(row=i, column=6, value=brk / 1440 if brk else 0)
    ws.cell(row=i, column=9, value=billable)

billable_dv = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
billable_dv.error = "Escribí Sí o No"
billable_dv.errorTitle = "Facturable"
ws.add_data_validation(billable_dv)

for r in range(2, ROWS + 2):
    # Horas = salida − entrada − descanso. MOD(…,1) hace que un turno que pasa
    # la medianoche (22:00 a 02:00) dé 4:00 y no un número negativo.
    ws.cell(row=r, column=7, value=f'=IF(OR(D{r}="",E{r}=""),"",MOD(E{r}-D{r},1)-F{r})')
    # Excel guarda las horas como fracción de día: ×24 las pasa a decimal.
    ws.cell(row=r, column=8, value=f'=IF(G{r}="","",ROUND(G{r}*24,2))')
    # Tarifa del cliente, desde la hoja Clientes. Se puede pisar a mano.
    ws.cell(
        row=r,
        column=10,
        value=f'=IF(B{r}="","",IFERROR(VLOOKUP(B{r},Clientes!$A$2:$B${CLIENT_ROWS + 1},2,FALSE),""))',
    )
    ws.cell(row=r, column=11, value=f'=IF(OR(H{r}="",J{r}=""),"",IF(I{r}="No",0,H{r}*J{r}))')

    ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
    ws.cell(row=r, column=4).number_format = "h:mm"
    ws.cell(row=r, column=5).number_format = "h:mm"
    ws.cell(row=r, column=6).number_format = "[h]:mm"
    ws.cell(row=r, column=7).number_format = "[h]:mm"
    ws.cell(row=r, column=8).number_format = "0.00"
    ws.cell(row=r, column=10).number_format = "#,##0.00"
    ws.cell(row=r, column=11).number_format = "#,##0.00"
    for c in (7, 8, 10, 11):
        ws.cell(row=r, column=c).fill = formula_fill
    for c in range(1, len(HEADERS) + 1):
        ws.cell(row=r, column=c).border = border
    billable_dv.add(ws.cell(row=r, column=9))

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:K{ROWS + 1}"
widths(ws, [12, 20, 44, 9, 9, 10, 9, 10, 11, 12, 12])

# ----------------------------------------------------------------------------
# Resumen
# ----------------------------------------------------------------------------
rs = wb.create_sheet("Resumen", 1)
rs["A1"] = "Resumen por cliente"
rs["A1"].font = title_font
rs["A2"] = "Se calcula solo a partir de las hojas Horas y Clientes. No hay que tocar nada acá."
rs["A2"].font = muted

rs.append([])
rs.append(["Cliente", "Horas", "Horas facturables", "Importe"])
style_header(rs, 4, 4)
first = 5
last = first + CLIENT_ROWS - 1
for r in range(first, last + 1):
    src = r - first + 2  # fila en Clientes
    rs.cell(row=r, column=1, value=f'=IF(Clientes!A{src}="","",Clientes!A{src})')
    rs.cell(row=r, column=2, value=f'=IF(A{r}="","",SUMIF(Horas!$B:$B,A{r},Horas!$G:$G))')
    # Facturable vacío cuenta como "Sí", igual que en la columna Importe de la
    # hoja Horas (que solo descuenta los "No"): si las dos columnas usaran
    # criterios distintos, una fila sin marcar sumaría plata pero no horas.
    rs.cell(row=r, column=3, value=f'=IF(A{r}="","",SUMIFS(Horas!$G:$G,Horas!$B:$B,A{r},Horas!$I:$I,"<>No"))')
    rs.cell(row=r, column=4, value=f'=IF(A{r}="","",SUMIF(Horas!$B:$B,A{r},Horas!$K:$K))')
    rs.cell(row=r, column=2).number_format = "[h]:mm"
    rs.cell(row=r, column=3).number_format = "[h]:mm"
    rs.cell(row=r, column=4).number_format = "#,##0.00"
    for c in range(1, 5):
        rs.cell(row=r, column=c).border = border

total_row = last + 1
rs.cell(row=total_row, column=1, value="Total")
rs.cell(row=total_row, column=2, value=f"=SUM(B{first}:B{last})")
rs.cell(row=total_row, column=3, value=f"=SUM(C{first}:C{last})")
rs.cell(row=total_row, column=4, value=f"=SUM(D{first}:D{last})")
rs.cell(row=total_row, column=2).number_format = "[h]:mm"
rs.cell(row=total_row, column=3).number_format = "[h]:mm"
rs.cell(row=total_row, column=4).number_format = "#,##0.00"
for c in range(1, 5):
    rs.cell(row=total_row, column=c).font = Font(bold=True)
    rs.cell(row=total_row, column=c).border = border
rs.cell(row=total_row + 1, column=1, value="El importe suma monedas distintas si tus clientes no cobran en la misma. En ese caso, mirá la fila de cada cliente.").font = muted

# Bloque por rango de fechas
rng = total_row + 3
rs.cell(row=rng, column=1, value="Horas entre dos fechas").font = title_font
rs.cell(row=rng + 1, column=1, value="Desde")
rs.cell(row=rng + 1, column=2, value=date(2026, 10, 19))
rs.cell(row=rng + 2, column=1, value="Hasta")
rs.cell(row=rng + 2, column=2, value=date(2026, 10, 25))
rs.cell(row=rng + 1, column=2).number_format = "DD/MM/YYYY"
rs.cell(row=rng + 2, column=2).number_format = "DD/MM/YYYY"
rs.cell(row=rng + 3, column=1, value="Horas")
rs.cell(row=rng + 3, column=2, value=f'=SUMIFS(Horas!$G:$G,Horas!$A:$A,">="&B{rng + 1},Horas!$A:$A,"<="&B{rng + 2})')
rs.cell(row=rng + 3, column=2).number_format = "[h]:mm"
rs.cell(row=rng + 4, column=1, value="Horas facturables")
rs.cell(
    row=rng + 4,
    column=2,
    value=f'=SUMIFS(Horas!$G:$G,Horas!$A:$A,">="&B{rng + 1},Horas!$A:$A,"<="&B{rng + 2},Horas!$I:$I,"<>No")',
)
rs.cell(row=rng + 4, column=2).number_format = "[h]:mm"
rs.cell(row=rng + 5, column=1, value="Importe")
rs.cell(row=rng + 5, column=2, value=f'=SUMIFS(Horas!$K:$K,Horas!$A:$A,">="&B{rng + 1},Horas!$A:$A,"<="&B{rng + 2})')
rs.cell(row=rng + 5, column=2).number_format = "#,##0.00"
for r in range(rng + 1, rng + 6):
    for c in (1, 2):
        rs.cell(row=r, column=c).border = border
rs.cell(row=rng + 6, column=1, value="Cambiá las dos fechas para ver una semana, un mes o el período de una factura.").font = muted
widths(rs, [26, 14, 18, 14])

# ----------------------------------------------------------------------------
# Instrucciones
# ----------------------------------------------------------------------------
ins = wb.create_sheet("Instrucciones")
LINES = [
    ("Planilla de control de horas trabajadas", title_font),
    ("De registruti.app/blog/control-de-horas-trabajadas. Gratis, para usar y modificar.", muted),
    ("", None),
    ("1. En la hoja Clientes, cargá cada cliente con su tarifa por hora y su moneda.", None),
    ("2. En la hoja Horas, una fila por bloque de trabajo: fecha, cliente, tarea, hora de entrada, hora de salida y descanso.", None),
    ("   El descanso va en formato de hora (0:45 son 45 minutos). Facturable: Sí o No (vacío cuenta como Sí).", None),
    ("3. Las columnas grises se calculan solas: horas en hh:mm, horas en decimal, tarifa del cliente e importe.", None),
    ("4. La hoja Resumen suma horas, horas facturables e importe por cliente, y también entre dos fechas.", None),
    ("", None),
    ("Para usarla en Google Sheets: subila a Google Drive, abrila con Google Sheets y elegí Archivo → Guardar como hoja de cálculo de Google. Las fórmulas funcionan igual.", None),
    ("", None),
    ("Las filas de ejemplo son la semana que cuenta el artículo. Borralas cuando empieces a cargar las tuyas (las fórmulas quedan).", None),
    ("", None),
    ("Cuando la planilla quede chica (varios clientes, monedas distintas, facturas cada mes), Registruti hace lo mismo gratis y genera el invoice en PDF desde las horas: https://registruti.app", None),
]
for i, (text, font) in enumerate(LINES, start=1):
    cell = ins.cell(row=i, column=1, value=text)
    if font:
        cell.font = font
    cell.alignment = Alignment(wrap_text=True, vertical="top")
ins.column_dimensions["A"].width = 110

wb.active = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"ok {OUT} ({OUT.stat().st_size} bytes)")
